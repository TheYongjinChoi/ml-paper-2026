from pathlib import Path
from collections import defaultdict, Counter
from html import escape
import json, re
from reportlab.pdfgen.canvas import Canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle, KeepTogether
from reportlab.platypus.tableofcontents import TableOfContents
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[1]
OUT=Path(r'C:\Users\chwou\Desktop\graduate\03_data\processed\서울시상권_코드북_2021_2025')
PDF=OUT/'변수코드북_읽기용.pdf'
D=json.loads((OUT/'codebook.json').read_text(encoding='utf-8'))
wb=load_workbook(OUT/'변수코드북.xlsx',read_only=True,data_only=False)
for sheet,key in [('변수코드북','variables'),('원자료열사전','raw'),('파일검증','files'),('업종코드','industries')]:
    actual=list(wb[sheet].iter_rows(min_row=2,values_only=True))
    assert len(actual)==len(D[key]),sheet
    for a,b in zip(actual,D[key]):
        assert [v or None for v in a]==[v or None for v in b.values()],(sheet,a,b)
pdfmetrics.registerFont(TTFont('Malgun','C:/Windows/Fonts/malgun.ttf'))
pdfmetrics.registerFont(TTFont('MalgunBold','C:/Windows/Fonts/malgunbd.ttf'))
pdfmetrics.registerFontFamily('Malgun',normal='Malgun',bold='MalgunBold',italic='Malgun',boldItalic='MalgunBold')
INK=colors.HexColor('#18364B'); MUTED=colors.HexColor('#536677')
PALE=colors.HexColor('#EEF3F7'); TEAL=colors.HexColor('#EAF4EF'); RULE=colors.HexColor('#CCD7DF')
W,H=595.276,841.89
M=43; WIDTH=W-2*M
styles={
 'body':ParagraphStyle('body',fontName='Malgun',fontSize=10,leading=16,wordWrap='CJK',spaceAfter=8,textColor=INK),
 'small':ParagraphStyle('small',fontName='Malgun',fontSize=8.5,leading=13,wordWrap='CJK',textColor=MUTED),
 'cell':ParagraphStyle('cell',fontName='Malgun',fontSize=9,leading=14,wordWrap='CJK',textColor=INK),
 'label':ParagraphStyle('label',fontName='MalgunBold',fontSize=9,leading=14,wordWrap='CJK',textColor=MUTED),
 'head':ParagraphStyle('head',fontName='MalgunBold',fontSize=18,leading=26,wordWrap='CJK',spaceAfter=14,textColor=INK,keepWithNext=True),
 'sub':ParagraphStyle('sub',fontName='MalgunBold',fontSize=13,leading=20,wordWrap='CJK',spaceBefore=10,spaceAfter=9,textColor=INK,keepWithNext=True),
 'var':ParagraphStyle('var',fontName='MalgunBold',fontSize=12,leading=18,wordWrap='CJK',spaceAfter=4,textColor=INK,keepWithNext=True),
 'cover':ParagraphStyle('cover',fontName='MalgunBold',fontSize=31,leading=43,wordWrap='CJK',spaceAfter=16,textColor=INK),
 'tablehead':ParagraphStyle('tablehead',fontName='MalgunBold',fontSize=9,leading=14,textColor=colors.white,wordWrap='CJK'),
}
def clean(s):
    return str(s).replace('\u2011','-').replace('−','-').replace('–','-').replace('—','-')
def p(s,style='body',markup=False):
    v=clean(s); return Paragraph(v if markup else escape(v).replace('\n','<br/>'),styles[style])
story=[]
bookmarks=[]
class Doc(SimpleDocTemplate):
    def afterFlowable(self,f):
        if hasattr(f,'bookmark'):
            key,title,level,toc=f.bookmark
            self.canv.bookmarkPage(key)
            self.canv.addOutlineEntry(title,key,level=level,closed=(level>0))
            if toc:self.notify('TOCEntry',(level,title,self.page,key))
            bookmarks.append((title,self.page,key))

def heading(title,key,level=0,toc=True,new=True):
    if new:story.append(PageBreak())
    v=p(title,'head' if level==0 else 'sub');v.bookmark=(key,title,level,toc);story.append(v)

def table(headers,rows,widths=None):
    widths=widths or [WIDTH/len(headers)]*len(headers)
    matrix=[[p(x,'tablehead') for x in headers]]+[[p(x,'cell') for x in row] for row in rows]
    t=Table(matrix,colWidths=widths,repeatRows=1,hAlign='LEFT')
    t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),INK),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),8),('RIGHTPADDING',(0,0),(-1,-1),8),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,PALE]),('LINEBELOW',(0,0),(-1,0),.6,RULE)]))
    return t

def card(r,number):
    title=p(f"{number:02d}  {r['변수']}",'var')
    title.bookmark=('v_'+r['변수코드'],r['변수코드']+' | '+r['변수'],1,False)
    meta=p(f"{r['변수코드']}  ·  {r['구분']}  ·  {r['구성요소']} / {r['특성군']}",'small')
    data=[
      ('범위 · 상태',r['분석범위']+' / '+r['채택상태']),
      ('정의 · 산식',r['조작적정의_산식']),
      ('원자료',r['원자료_열']),
      ('단위 · 시간',r['단위']+' / '+r['시간배치']),
      ('자료형 · 범위',r['자료형_R']+' / '+r['허용범위']),
      ('개념 연결',r['개념연결_예상관계'] or '분석을 위한 식별·결합 정보'),
      ('결측 처리',r['결측_처리']),
      ('해석상 주의',r['해석상주의'] or '공통 분석·결합 원칙 적용'),
      ('근거',r['근거'])]
    t=Table([[p(a,'label'),p(b,'cell')] for a,b in data],colWidths=[76,WIDTH-76],hAlign='LEFT')
    t.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('BACKGROUND',(0,0),(0,-1),PALE),('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4),('LINEABOVE',(0,0),(-1,0),.8,RULE),('LINEBELOW',(0,-1),(-1,-1),.8,RULE)]))
    # A variable stays together so its formula and cautions are visible on one page.
    story.append(KeepTogether([title,meta,Spacer(1,6),t,Spacer(1,17)]))

story += [Spacer(1,70),p('서울시 상권분석서비스','sub'),p('변수 코드북\n2021-2025','cover'),p('지역사회 경제 리질리언스 연구를 위한\n개념·지표·원자료 연결','body'),Spacer(1,28)]
story.append(table(['주분석','보조분석'],[['골목상권 · 발달상권 · 전통시장','골목상권의 상권배후지'],['상권의 경제적 성과와 내부·외부 설명요인','가외성·내구성 관련 조건과 경쟁 조건']], [WIDTH/2]*2))
story += [Spacer(1,28),p('83개 변수·변수군  /  453개 표준화 원자료 열\n35개 CSV 파일 검증  /  100개 생활밀접업종','body'),Spacer(1,36),p('작성 기준  2026.10.11\n기존 엑셀 코드북을 읽기용으로 재구성. 측정 설계안이며 실증분석 결과가 아님.','small')]
heading('목차','contents')
toc=TableOfContents()
toc.levelStyles=[ParagraphStyle('toc0',fontName='Malgun',fontSize=11,leading=18,spaceBefore=8,textColor=INK),ParagraphStyle('toc1',fontName='Malgun',fontSize=9,leading=14,leftIndent=14,textColor=MUTED)]
story += [p('목차와 PDF 책갈피에서 원하는 부분으로 이동할 수 있습니다. 변수별 책갈피에는 영문 코드와 한글 이름을 함께 표시했습니다.','small'),Spacer(1,16),toc]

heading('01  읽는 방법과 분석 범위','guide')
story += [p('이 문서는 경제적 성과, 리질리언스 구성요소에 연결한 설명요인, 통제변수, 배후지 보조분석을 구분합니다. 표의 모든 후보를 하나의 모형에 동시에 넣는다는 뜻은 아닙니다.'),
table(['구분','적용 범위'],[['주분석','2021Q1-2025Q4. 골목상권(A)·발달상권(D)·전통시장(R). 관광특구(U) 제외.'],['보유 상권 범위','점포·영역 기준 총 1,644개: 골목 1,090개, 발달 249개, 전통시장 305개. 최종 분석 표본은 성과별 관측 범위에 따라 달라짐.'],['배후지 보조','보유 통계는 골목상권 배후지 1,090개. 해당 상권의 성과를 설명하는 보조자료로 사용. 배후지를 별도의 주분석 community로 합치지 않음.'],['자료 한계','조직·금융·외부 연대성은 추가 기록이 필요. 배후지 시설은 2025년 4개 분기만 보유. 배후지 영역은 1,071개로 통계 코드와 차이.']],[90,WIDTH-90]),Spacer(1,15),
p('변수 상태를 읽는 방법','sub'),table(['표시','의미'],[['산출 가능·비교성 검증 필요','원자료 열은 보유. 공간·시계열·모집단 정합성을 확인해야 최종 채택 가능.'],['보조 후보 / 기술통계 보조','주 구성요소의 직접 측정이나 주 분석 설명변수로 확정하지 않은 항목.'],['추가 자료 필요 / 투입 보류','미확보 기록 또는 정의·시점 문제가 있어 현재 분석 입력에 사용하지 않음.'],['분류표 확정 후 산출','기능군 또는 정책 제한 업종의 대응표를 먼저 확정해야 함.']],[165,WIDTH-165])]
heading('02  구성요소별 분류와 공통 원칙','principles')
story.append(p('6Rs는 하현상 외(2014)의 이론적 구분입니다. 개별 상권 지표와의 연결은 본 연구의 조작화이며 원문의 검증된 척도를 그대로 옮긴 것이 아닙니다. 민감성·대응력은 6Rs와 별개의 분석 관점으로 유지합니다.'))
for r in D['principles']:
    story.append(KeepTogether([p(r['항목'],'sub'),p(r['적용원칙'])]))

heading('03  변수 빠른 찾기','index')
idxrows=[]
for i,r in enumerate(D['variables'],1):
    name=Paragraph(f'<link href="#v_{r["변수코드"]}" color="#18364B">{escape(r["변수"])}</link>',styles['cell'])
    idxrows.append([p(str(i),'cell'),p(r['변수코드'],'cell'),name,p(r['구분']+'\n'+r['구성요소'],'cell')])
t=Table([[p(x,'tablehead') for x in ['번호','변수 코드','변수명 (클릭하여 이동)','분류']]]+idxrows,colWidths=[30,151,177,WIDTH-358],repeatRows=1)
t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),INK),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,PALE]),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]));story.append(t)

sections=[('04  식별자와 경제적 성과','outcomes',lambda r:r['구분'] in ['식별자','종속변수','보조 결과'] and '배후지' not in r['분석범위']),
          ('05  리질리언스 구성요소별 설명·과정 변수','explanatory',lambda r:r['구분'] in ['설명변수','과정변수'] and '배후지' not in r['분석범위']),
          ('06  통제변수와 과거 실적','controls',lambda r:r['구분'] in ['통제변수','예측 통제'] and '배후지' not in r['분석범위']),
          ('07  상권배후지 보조분석','hinterland',lambda r:'배후지' in r['분석범위']),
          ('08  내부·외부 조건의 상호작용','interaction',lambda r:r['구분']=='상호작용')]
seen=[]
for title,key,predicate in sections:
    heading(title,key)
    if key=='explanatory':story.append(p('내구성·가외성·자원부존성·신속성·적응성·연대성과의 연결을 표시합니다. 결과의 유지·재편과 원인 후보를 구분하고, 추가 자료가 필요한 항목은 현재 확보된 변수와 구별합니다.'))
    if key=='hinterland':story.append(p('배후지 경쟁 조건은 별도의 6Rs 구성요소로 취급하지 않습니다. 같은 공급이 대체 가능성과 경쟁 압력 중 어떤 관계를 보이는지 비교합니다. 주분석 상권의 성과를 유지한 보조모형으로 검토합니다.'))
    for i,r in enumerate(D['variables'],1):
        if predicate(r):card(r,i);seen.append(r['변수코드'])
assert len(seen)==83 and len(set(seen))==83

heading('부록 A  원자료 열 대응표','raw')
story.append(p('공간·자료군별 표준화 열 453개를 정리했습니다. 원래의 한글·영문 열 이름은 보존하며, 오기 후보나 미확인 정의는 주의사항에 남겼습니다. F번호는 부록 B의 원본 파일을 가리킵니다.'))
groups=defaultdict(list)
for r in D['raw']:groups[(r['자료군'],r['공간'])].append(r)
for j,((fam,scope),items) in enumerate(groups.items()):
    heading(f'{fam} - {scope}',f'raw{j}',1,False,new=False)
    files=sorted(set(';'.join(r['출처파일'] for r in items).split(';')))
    periods=sorted(set(r['관측분기수'] for r in items))
    story.append(p('파일: '+', '.join(files)+' / 관측 분기 수: '+', '.join(map(str,periods))+' / 실제 열 이름이 표준 열 이름과 같으면 ‘동일’로 표시.','small'))
    special=list(dict.fromkeys(r['주의'] for r in items if not r['주의'].startswith('열 이름은 원문 보존')))
    for note in special:story.append(p(note,'small'))
    content=[]
    for r in items:
        actual=r['실제열명']
        if actual==r['표준열명']:actual='동일'
        else:actual=actual.replace(r['표준열명']+' / ','').replace(' / '+r['표준열명'],'')
        content.append([r['표준열명'],actual,r['단위']])
    story.append(table(['표준 열 이름','다른 실제 원자료 열 이름','단위'],content,[182,210,WIDTH-392]))
    story.append(Spacer(1,12))

heading('부록 B  원본 파일과 자료 검증','files')
story.append(p('35개 CSV를 직접 읽어 2021-2025년과 대상 유형을 필터링하고 후보 키를 확인했습니다. 모든 파일에서 후보 키 중복 초과행은 0이었습니다. 이는 값 전체의 타당성이나 공간 기준 일치가 검증됐다는 뜻은 아닙니다.'))
story.append(p('원자료 기준 폴더: 03_data/raw/서울시 상권분석서비스. 파일명·자료군으로 원본을 찾을 수 있으며, 엑셀 코드북의 파일검증 시트에는 절대경로가 수록되어 있습니다.','small'))
for r in D['files']:
    story.append(KeepTogether([p(f"{r['자료ID']}  {r['자료군']} - {r['공간']}",'sub'),p(r['파일명'],'cell'),p(f"기간: {r['최초분기']}~{r['최종분기']} ({r['관측분기수']}분기) / 인코딩: {r['인코딩']}\n전체 {r['전체행']:,}행 → 기간 내 {r['기간내행']:,}행 → 대상 {r['대상행']:,}행\n대상 상권 {r['대상상권수']:,}개 / 후보 키 중복 {r['키중복']}행 / 유형별 상권 수: {r['유형']}",'cell'),Spacer(1,8)]))

heading('부록 C  생활밀접업종 코드','industries')
story.append(p('실제 관측한 100개 업종입니다. 대분류는 코드 접두부를 따릅니다. 기능군과 방역정책별 제한 분류는 아직 확정되지 않았으므로 임의로 채우지 않았습니다.'))
story.append(table(['업종 코드','업종명','대분류'],[[r['업종코드'],r['업종명'],r['대분류']] for r in D['industries']],[110,285,WIDTH-395]))
heading('부록 D  이론·자료 출처','sources')
for r in D['sources']:
    title=f"{r['ID']}  {r['저자_기관']}"
    body=[p(title,'sub'),p(r['근거'])]
    if r['위치'].startswith('http'):
        body.append(p(f'<link href="{escape(r["위치"],quote=True)}" color="#245B80">{escape(r["위치"])}</link>','small',True))
    else:body.append(p(r['위치'],'small'))
    story.append(KeepTogether(body+[Spacer(1,8)]))

def page(c,doc):
    c.saveState()
    if doc.page>1:
        c.setFont('Malgun',8);c.setFillColor(MUTED)
        c.drawString(M,H-27,'서울시 상권분석서비스  |  변수 코드북 2021-2025')
        c.setStrokeColor(RULE);c.setLineWidth(.5);c.line(M,H-34,W-M,H-34)
    c.setFont('Malgun',8);c.setFillColor(MUTED)
    c.drawString(M,25,'측정 설계안 · 2026.10.11')
    c.drawRightString(W-M,25,str(doc.page))
    c.restoreState()

doc=Doc(str(PDF),pagesize=(W,H),rightMargin=M,leftMargin=M,topMargin=48,bottomMargin=45,title='서울시 상권분석서비스 변수 코드북 2021-2025',author='연구자',subject='리질리언스 구성요소별 측정 설계와 원자료 사전')
doc.multiBuild(story,onFirstPage=page,onLaterPages=page)
reader=PdfReader(PDF)
text='\n'.join(x.extract_text() or '' for x in reader.pages)
missing=[r['변수코드'] for r in D['variables'] if r['변수코드'] not in text]
assert not missing,missing
report={'pages':len(reader.pages),'variable_count':len(seen),'raw_fields':len(D['raw']),'missing_variable_codes':missing,'file':str(PDF),'bookmarks':bookmarks[-100:]}
(ROOT/'codebook'/'pdf_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k!='bookmarks'},ensure_ascii=False))
